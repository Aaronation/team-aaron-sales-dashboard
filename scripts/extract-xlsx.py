#!/usr/bin/env python3
"""Extract Team Aaron Sales Data sheets into sales-line-items.json."""

from __future__ import annotations

import argparse
import json
import os
from collections import Counter
from datetime import date, datetime

from openpyxl import load_workbook


def parse_date(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if v is None:
        return None
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s[:10], fmt).date().isoformat()
        except Exception:
            pass
    return None


def extract(path: str, source: str):
    wb = load_workbook(path, data_only=False, read_only=True)
    ws = wb["Data"]
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter)
    headers = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(header)]
    idx = {h: i for i, h in enumerate(headers)}

    def g(row, *names, default=None):
        for n in names:
            if n in idx and idx[n] < len(row):
                return row[idx[n]]
        return default

    out = []
    months = Counter()
    salespeople = Counter()
    depts = Counter()
    gal_lines = 0

    for row in rows_iter:
        if not row or all(c is None for c in row):
            continue
        inv = g(row, "Invoice")
        if inv is None:
            continue
        d = parse_date(g(row, "Date"))
        year = month = week = dayname = None
        if d:
            dt = datetime.fromisoformat(d)
            year = str(dt.year)
            month = dt.strftime("%B")
            week = int(dt.strftime("%U"))
            dayname = dt.strftime("%A")
            months[d[:7]] += 1

        qty = g(row, "Quantity") or 0
        unit = g(row, "Unit Price") or 0
        total = g(row, "Total Amount")
        try:
            qty = float(qty)
            unit = float(unit)
            total = float(total) if total is not None else qty * unit
        except Exception:
            continue

        sp = g(row, "Salesperson")
        dept = g(row, "Sales Department")
        if sp:
            salespeople[str(sp)] += 1
        if dept:
            depts[str(dept)] += 1
        inv_s = str(inv)
        if inv_s.upper().startswith("GAL"):
            gal_lines += 1

        out.append(
            {
                "source": source,
                "invoice": inv_s,
                "tradingPartner": str(g(row, "Trading Partner")) if g(row, "Trading Partner") else None,
                "customer": str(g(row, "Customer Name")) if g(row, "Customer Name") else None,
                "date": d,
                "year": year,
                "month": month,
                "week": week,
                "day": dayname,
                "department": str(dept) if dept else None,
                "salesperson": str(sp) if sp else None,
                "sku": str(g(row, "SKU")) if g(row, "SKU") else None,
                "productModel": str(g(row, "产品型号")) if g(row, "产品型号") else None,
                "category": str(g(row, "Name")) if g(row, "Name") else None,
                "quantity": qty,
                "unitPrice": unit,
                "totalAmount": total,
                "productName": str(g(row, "Product Name"))[:120] if g(row, "Product Name") else None,
                "customerCode": str(g(row, "Customer code")) if g(row, "Customer code") else None,
            }
        )

    wb.close()
    meta = {
        "source": source,
        "rows": len(out),
        "months": dict(months.most_common()),
        "salespeople": dict(salespeople),
        "departments": dict(depts),
        "galInvoicesLines": gal_lines,
        "dateMin": min((r["date"] for r in out if r["date"]), default=None),
        "dateMax": max((r["date"] for r in out if r["date"]), default=None),
        "revenue": sum(r["totalAmount"] for r in out),
    }
    return out, meta


def row_key(r):
    return (r["invoice"], r["sku"], r["date"], r["quantity"], r["totalAmount"], r["customer"])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("oct_xlsx")
    ap.add_argument("aug_xlsx")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    oct_rows, oct_meta = extract(args.oct_xlsx, "202610")
    aug_rows, aug_meta = extract(args.aug_xlsx, "202608")
    keys = {row_key(r) for r in oct_rows}
    merged = list(oct_rows)
    for r in aug_rows:
        if row_key(r) not in keys:
            merged.append(r)

    payload = {
        "meta": {
            "generatedFrom": [os.path.basename(args.oct_xlsx), os.path.basename(args.aug_xlsx)],
            "primaryPeriod": "2026-10",
            "oct": oct_meta,
            "aug": aug_meta,
            "mergedRows": len(merged),
            "targets": {
                "annualTeamTarget": 3500000,
                "year": 2026,
                "splits": {"Aaron Chang": 0.49, "Kelvin Lai": 0.31, "Kenny Ma": 0.2},
            },
        },
        "lineItems": merged,
    }

    os.makedirs(os.path.dirname(os.path.abspath(args.out)) or ".", exist_ok=True)
    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, separators=(",", ":"))
    print(f"Wrote {args.out} ({len(merged)} rows)")


if __name__ == "__main__":
    main()
